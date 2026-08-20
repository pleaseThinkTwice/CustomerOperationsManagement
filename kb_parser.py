# -*- coding: utf-8 -*-
"""
知识库文件解析器：把 pdf / docx / pptx / xlsx / html / txt 解析成纯文本。

所有第三方库都做惰性导入 + 优雅降级：缺库时该格式返回明确错误而非崩溃。
返回统一结构：{"ok": bool, "text": str, "error": str}
"""
import os

SUPPORTED_EXT = {".pdf", ".docx", ".pptx", ".xlsx", ".xls", ".html", ".htm", ".txt", ".md", ".csv"}


def _ok(text):
    return {"ok": True, "text": text or "", "error": ""}


def _err(msg):
    return {"ok": False, "text": "", "error": msg}


def parse_txt(path):
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return _ok(f.read())
    except Exception as e:
        return _err(f"txt 读取失败: {e}")


def parse_pdf(path):
    try:
        from pypdf import PdfReader
    except Exception:
        return _err("未安装 pypdf，无法解析 pdf（pip install pypdf）")
    try:
        reader = PdfReader(path)
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        return _ok("\n".join(parts))
    except Exception as e:
        return _err(f"pdf 解析失败: {e}")


def parse_docx(path):
    try:
        import docx  # python-docx
    except Exception:
        return _err("未安装 python-docx，无法解析 docx（pip install python-docx）")
    try:
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        # 表格
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return _ok("\n".join(parts))
    except Exception as e:
        return _err(f"docx 解析失败: {e}")


def parse_pptx(path):
    try:
        from pptx import Presentation
    except Exception:
        return _err("未安装 python-pptx，无法解析 pptx（pip install python-pptx）")
    try:
        prs = Presentation(path)
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"[幻灯片 {i}]")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = "".join(run.text for run in para.runs)
                        if t.strip():
                            parts.append(t)
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        if any(cells):
                            parts.append(" | ".join(cells))
        return _ok("\n".join(parts))
    except Exception as e:
        return _err(f"pptx 解析失败: {e}")


def parse_xlsx(path):
    try:
        import openpyxl
    except Exception:
        return _err("未安装 openpyxl，无法解析 xlsx（pip install openpyxl）")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[工作表 {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(c.strip() for c in cells):
                    parts.append(" | ".join(cells))
        wb.close()
        return _ok("\n".join(parts))
    except Exception as e:
        return _err(f"xlsx 解析失败: {e}")


def parse_html(path):
    try:
        from bs4 import BeautifulSoup
    except Exception:
        return _err("未安装 beautifulsoup4，无法解析 html（pip install beautifulsoup4）")
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            soup = BeautifulSoup(f.read(), "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        return _ok(soup.get_text("\n", strip=True))
    except Exception as e:
        return _err(f"html 解析失败: {e}")


_DISPATCH = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".pptx": parse_pptx,
    ".xlsx": parse_xlsx,
    ".xls": parse_xlsx,
    ".html": parse_html,
    ".htm": parse_html,
    ".txt": parse_txt,
    ".md": parse_txt,
    ".csv": parse_txt,
}


def parse_file(path):
    """按扩展名分派。返回 {ok, text, error}。"""
    ext = os.path.splitext(path)[1].lower()
    fn = _DISPATCH.get(ext)
    if not fn:
        return _err(f"不支持的文件类型: {ext}")
    return fn(path)


def chunk_text(text, max_chars=1200, overlap=100):
    """把长文本切成较小片段（按段落聚合，避免单片过大）。"""
    text = (text or "").strip()
    if not text:
        return []
    paras = [p.strip() for p in text.split("\n") if p.strip()]
    chunks, buf = [], ""
    for p in paras:
        if len(buf) + len(p) + 1 <= max_chars:
            buf = (buf + "\n" + p) if buf else p
        else:
            if buf:
                chunks.append(buf)
            # 超长单段直接硬切
            while len(p) > max_chars:
                chunks.append(p[:max_chars])
                p = p[max_chars - overlap:]
            buf = p
    if buf:
        chunks.append(buf)
    return chunks
